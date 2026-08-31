"""
Wood Pole Replacement — constraint matrix tool for ArcGIS Pro.

2026 update:
  * 2026 streaming structure layer field names (StructureN / Latitude / Longitude),
    plus Height1, Installati and Structure T carried into the matrix.
  * Constraint layers report the intersecting feature's name instead of "Yes",
    resolved by an override -> known alias -> label expression -> heuristic ladder.
  * Named-layer analyses: county, township/city, section-township-range,
    MDOT jurisdiction, soil, parcel.
  * The 100 ft buffer is replaced by a 500 ft buffer used solely for the SESC
    check, against a per-county choice of water / wetland / county drain layers.
    Every other overlap stays a question about the 60 ft work zone.
  * Optional project-wide 2 mile MNFI T&E review.

Toolbox parameters are listed in docs/pole-tool-2026-update-plan.md, section 9.
"""

import arcpy
import os
import re
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl import load_workbook

arcpy.env.overwriteOutput = True

# --- Inputs ---

poles = arcpy.GetParameter(0)
access_routes = arcpy.GetParameter(1)
constraint_layers = arcpy.GetParameter(2) or []
out_folder = arcpy.GetParameterAsText(3)
project_name = arcpy.GetParameterAsText(4)
pole_id_field = arcpy.GetParameterAsText(5)
ar_pole_id_field = arcpy.GetParameterAsText(
    6
)  # Field in access routes with comma-separated pole IDs

county_layer = arcpy.GetParameter(7)
township_layer = arcpy.GetParameter(8)
plss_layer = arcpy.GetParameter(9)
mdot_layer = arcpy.GetParameter(10)
soil_layer = arcpy.GetParameter(11)
parcel_layer = arcpy.GetParameter(12)
sesc_layers = arcpy.GetParameter(13) or []
sesc_from_work_zone = bool(arcpy.GetParameter(14))
run_te_review = bool(arcpy.GetParameter(15))
mnfi_layer = arcpy.GetParameter(16)
field_override_param = arcpy.GetParameter(17)

gdb = arcpy.env.scratchGDB

# Joins the values of a multi-field answer, e.g. "MdB - Miami loam".
VALUE_SEP = " - "

RUN_LOG = []
_temp_counter = [0]


def log(category, subject, message):
    """Record a decision or problem for the tool messages and the Run Log sheet."""
    RUN_LOG.append({"Category": category, "Layer": subject, "Detail": message})
    arcpy.AddMessage(f"[{category}] {subject}: {message}")


def temp_name(prefix):
    """A scratch name that will not collide with anything already in the session."""
    _temp_counter[0] += 1
    return f"{prefix}_{_temp_counter[0]}"


def layer_title(layer):
    """A readable name for a layer, whatever form the parameter arrived in."""
    return str(getattr(layer, "name", layer))


def feature_count(source):
    """How many features a layer holds, or None if it cannot be counted."""
    try:
        return int(arcpy.management.GetCount(source)[0])
    except Exception:
        return None


# ---------------------------------------------------------------------------
# Field resolution
#
# County-specific layers name the same thing differently every time — one
# county's drain name lives in "Label", another's in "County Drain Name". These
# helpers work out which field holds the name we want to report, and record how
# they decided so a wrong guess is visible instead of silent.
# ---------------------------------------------------------------------------


def _norm(value):
    """Collapse a field name or alias to a comparable form: lowercase alphanumerics."""
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def field_index(source):
    """[(name, alias, type)] for a layer or feature class, minus geometry and ID fields."""
    fields = []
    for f in arcpy.ListFields(source):
        if f.type in ("OID", "Geometry", "GlobalID", "Raster", "Blob"):
            continue
        if f.name.lower() in ("shape", "shape_length", "shape_area"):
            continue
        fields.append((f.name, f.aliasName or f.name, f.type))
    return fields


def find_field(fields, candidates):
    """First field whose name or alias matches one of `candidates`, in candidate order."""
    by_name = {}
    by_alias = {}
    for name, alias, _type in fields:
        by_name.setdefault(_norm(name), name)
        by_alias.setdefault(_norm(alias), name)
    for candidate in candidates:
        key = _norm(candidate)
        if key in by_name:
            return by_name[key]
        if key in by_alias:
            return by_alias[key]
    return None


# Every real-world way a label expression names a field. The originals only
# handled \w+, which silently failed on any field with a space in it — that is
# why layers like MDOT's "Legal System" reported "Yes" instead of a name.
LABEL_PATTERNS = [
    r'\$feature\.([A-Za-z_][A-Za-z0-9_]*)',   # $feature.NAME
    r'\$feature\[\s*"([^"]+)"\s*\]',          # $feature["Name With Spaces"]
    r"\$feature\[\s*'([^']+)'\s*\]",          # $feature['Name With Spaces']
    r'\[([^\[\]]+)\]',                        # [Name With Spaces]
]


def parse_label_fields(layer):
    """Field names referenced by a layer's label expressions, in the order they appear."""
    found = []
    try:
        label_classes = list(layer.listLabelClasses())
    except Exception:
        # Streaming layers routinely refuse this; it is not an error, just a dead end.
        return found

    # Visible label classes first — those are the ones the cartographer meant.
    label_classes.sort(key=lambda lc: not getattr(lc, "visible", True))

    for label_class in label_classes:
        expression = getattr(label_class, "expression", "") or ""
        for pattern in LABEL_PATTERNS:
            for match in re.finditer(pattern, expression):
                name = match.group(1).strip().strip("\"'")
                if name and name not in found:
                    found.append(name)
    return found


# Keywords that suggest a field holds the name we want, per kind of layer.
# Ordered most specific first so infer_role() picks "drain" over "county" for a
# layer called "Kent County Drains".
ROLE_KEYWORDS = {
    "drain": ["drain", "ditch", "drainname"],
    "soil": ["musym", "muname", "mapunit", "soil"],
    "parcel": ["parcel", "pin", "taxid", "owner", "propertyid"],
    "plss": ["strlabel", "section", "township", "range"],
    "road": ["legalsystem", "jurisdiction", "roadname", "route", "nfc"],
    "water": ["water", "stream", "lake", "pond", "river", "wetland", "gnis"],
    "township": ["twp", "municipality", "city", "village"],
    "county": ["county", "cnty"],
    "te": ["sname", "comname", "ename", "element", "species", "status"],
}

STRONG_NAMES = ("label", "strlabel", "name", "musym", "muname", "legalsystem")
NEGATIVE_NAMES = ("objectid", "fid", "globalid", "shape", "hyperlink", "created", "edited", "symbol")


def infer_role(name):
    """Guess what kind of layer this is from its name, to bias field scoring."""
    text = _norm(name)
    for role, keywords in ROLE_KEYWORDS.items():
        if role in text or any(keyword in text for keyword in keywords):
            return role
    return None


def score_field(name, alias, ftype, role):
    """How likely a field is to hold the human-readable name for this kind of layer."""
    if any(bad in _norm(name) for bad in NEGATIVE_NAMES):
        return -1

    # Numbers and dates are rarely the name, but a coded drain number beats nothing.
    score = -5 if ftype in ("Double", "Integer", "SmallInteger", "Single", "Date") else 0

    for form in (_norm(name), _norm(alias)):
        if form in STRONG_NAMES:
            score += 100
        if form.endswith("name") or form.startswith("name"):
            score += 40
        if "label" in form:
            score += 35
        elif "name" in form:
            score += 25
        if "desc" in form:
            score += 10
        for keyword in ROLE_KEYWORDS.get(role, []):
            if keyword in form:
                score += 30
    return score


def parse_overrides(param):
    """{normalized layer name: [field names]} from the Field Overrides value table."""
    overrides = {}
    if not param:
        return overrides

    rows = []
    try:
        for r in range(param.rowCount):
            rows.append((param.getValue(r, 0), param.getValue(r, 1)))
    except AttributeError:
        for line in str(param).splitlines():
            parts = line.split(None, 1)
            if len(parts) == 2:
                rows.append(parts)

    for name, field_text in rows:
        if not name or not field_text:
            continue
        fields = [f.strip() for f in str(field_text).split(",") if f.strip()]
        if fields:
            overrides[_norm(str(name).strip("'\""))] = fields
    return overrides


def resolve_fields(layer, title, source, role=None, preferred=None, want=1, overrides=None):
    """
    Decide which field(s) hold the name to report for a layer.

    Ladder, first hit wins: an explicit override from the Field Overrides
    parameter, then the field names/aliases we already know for this kind of
    analysis, then the layer's own label expression, then keyword scoring.
    Returns (field names, provenance) — provenance goes in the Run Log so a bad
    guess can be corrected with an override instead of a code change.
    """
    fields = field_index(source)
    if not fields:
        return [], "layer has no usable attribute fields"

    override = (overrides or {}).get(_norm(title))
    if override:
        picked = [f for f in (find_field(fields, [c]) for c in override) if f]
        if picked:
            return picked, f"override -> {', '.join(picked)}"
        log("Field", title, f"override {override} matched no field on the layer, falling back")

    if preferred:
        picked = []
        for group in preferred:
            hit = find_field(fields, group if isinstance(group, list) else [group])
            if hit and hit not in picked:
                picked.append(hit)
        if len(picked) >= want:
            return picked[:want], f"known field/alias -> {', '.join(picked[:want])}"

    for candidate in parse_label_fields(layer):
        hit = find_field(fields, [candidate])
        if hit:
            return [hit], f"label expression -> {hit}"

    ranked = sorted(
        ((score_field(n, a, t, role), n) for n, a, t in fields),
        key=lambda pair: -pair[0],
    )
    picked = [name for score, name in ranked[:want] if score > 0]
    if picked:
        return picked, f"best guess -> {', '.join(picked)} (override if wrong)"

    available = ", ".join(name for name, _a, _t in fields[:25])
    return [], f"no name field found; reporting Yes/No. Fields available: {available}"


# ---------------------------------------------------------------------------
# Layer access
# ---------------------------------------------------------------------------


def ensure_feature_layer(layer, scratch_gdb, aoi_features=None, aoi_extent_fc=None):
    desc = arcpy.Describe(layer)

    # Detect web/streaming layers — these have shapeType but don't
    # work reliably with geoprocessing tools like SpatialJoin
    is_web = getattr(layer, "isWebLayer", False)
    if not is_web and hasattr(desc, "catalogPath"):
        is_web = str(desc.catalogPath).startswith("http")
    if not is_web and hasattr(desc, "path"):
        is_web = str(desc.path).startswith("http")

    # Local feature layer/class — safe to use directly
    if not is_web and hasattr(desc, "shapeType"):
        return layer

    # Web/service/streaming layer — three strategies, each falling back to the next:
    #   1. Select by actual project features (most precise)
    #   2. Select by simple AOI rectangle (simpler query)
    #   3. CopyFeatures with extent environment (last resort)
    name = layer_title(layer)
    arcpy.AddMessage(f"Copying web layer '{name}' locally...")
    output_fc = arcpy.CreateUniqueName("web_layer", scratch_gdb)

    def _try_select_and_copy(select_features):
        sel = temp_name("web_download_sel")
        try:
            arcpy.management.MakeFeatureLayer(layer, sel)
            for i, aoi_fc in enumerate(select_features):
                arcpy.management.SelectLayerByLocation(
                    sel, "INTERSECT", aoi_fc,
                    selection_type="NEW_SELECTION" if i == 0 else "ADD_TO_SELECTION",
                )
            copied = arcpy.management.CopyFeatures(sel, output_fc)[0]
            if feature_count(copied) == 0:
                raise ValueError("selection returned 0 features")
            return copied
        finally:
            if arcpy.Exists(sel):
                arcpy.management.Delete(sel)

    # Strategy 1: Select by actual project features
    if aoi_features:
        try:
            temp_fc = _try_select_and_copy(aoi_features)
            arcpy.AddMessage("  Downloaded via project feature selection.")
        except Exception as exc:
            arcpy.AddMessage(f"  Selection by project features failed ({exc}), trying next method...")
            temp_fc = None
        if temp_fc:
            return arcpy.management.MakeFeatureLayer(temp_fc, temp_name("temp_lyr"))[0]

    # Strategy 2: Select by simple AOI rectangle
    if aoi_extent_fc:
        try:
            temp_fc = _try_select_and_copy([aoi_extent_fc])
            arcpy.AddMessage("  Downloaded via AOI rectangle selection.")
        except Exception as exc:
            arcpy.AddMessage(f"  Selection by AOI rectangle failed ({exc}), trying next method...")
            temp_fc = None
        if temp_fc:
            return arcpy.management.MakeFeatureLayer(temp_fc, temp_name("temp_lyr"))[0]

    # Strategy 3: CopyFeatures with extent environment (last resort)
    arcpy.AddMessage(f"  Falling back to extent-based download for '{name}'...")
    original_extent = arcpy.env.extent
    if aoi_extent_fc:
        arcpy.env.extent = arcpy.Describe(aoi_extent_fc).extent
    try:
        temp_fc = arcpy.management.CopyFeatures(layer, output_fc)[0]
    finally:
        arcpy.env.extent = original_extent

    return arcpy.management.MakeFeatureLayer(temp_fc, temp_name("temp_lyr"))[0]


# ---------------------------------------------------------------------------
# Intersection
# ---------------------------------------------------------------------------


def _parse_ids(raw_id, multi_id):
    """Pole IDs as stripped strings, optionally splitting a comma-separated list."""
    if raw_id is None:
        return []
    text = str(raw_id).strip()
    if not text:
        return []
    if multi_id:
        return [part.strip() for part in text.split(",") if part.strip()]
    return [text]


def _resolve_sj_field(sj_fields, wanted, prefer_suffixed):
    """
    Find `wanted` in a SpatialJoin output, where colliding names get a numeric suffix.

    When both inputs carry the same field name the join layer's copy is the
    suffixed one, so `prefer_suffixed` says which side we are after.
    """
    if not prefer_suffixed and wanted in sj_fields:
        return wanted

    matches = [f for f in sj_fields if _norm(re.sub(r"_\d+$", "", f)) == _norm(wanted)]
    if not matches:
        return wanted if wanted in sj_fields else None

    matches.sort(key=lambda f: 0 if bool(re.search(r"_\d+$", f)) == prefer_suffixed else 1)
    return matches[0]


def get_intersect_hits(target, join_layer, id_field, label_fields, scratch_gdb, multi_id=False):
    """
    {pole id: set of value tuples} for features of `join_layer` intersecting `target`.

    Each tuple holds one value per entry in `label_fields`, so callers can report
    them together or in separate columns. Without `label_fields` the tuple is
    ("Yes",).
    """
    hits = {}

    if label_fields:
        sj_out = arcpy.CreateUniqueName("sj_temp", scratch_gdb)
        try:
            arcpy.analysis.SpatialJoin(
                target, join_layer, sj_out,
                join_operation="JOIN_ONE_TO_MANY",
                join_type="KEEP_COMMON",
                match_option="INTERSECT",
            )

            sj_fields = [f.name for f in arcpy.ListFields(sj_out)]
            id_col = _resolve_sj_field(sj_fields, id_field, prefer_suffixed=False)
            value_cols = [
                _resolve_sj_field(sj_fields, field, prefer_suffixed=True)
                for field in label_fields
            ]
            if not id_col or not all(value_cols):
                log("Warning", layer_title(join_layer),
                    "name field vanished from the spatial join output; reporting Yes/No")
                return get_intersect_hits(target, join_layer, id_field, None, scratch_gdb, multi_id)

            with arcpy.da.SearchCursor(sj_out, [id_col] + value_cols) as cursor:
                for row in cursor:
                    values = tuple(
                        "" if v is None else str(v).strip() for v in row[1:]
                    )
                    for pid in _parse_ids(row[0], multi_id):
                        # A feature with a blank name still intersected; never
                        # let a missing attribute turn a hit into a "No".
                        hits.setdefault(pid, set()).add(values if any(values) else ("Yes",))
        finally:
            if arcpy.Exists(sj_out):
                arcpy.management.Delete(sj_out)
    else:
        sel = temp_name("intersect_sel")
        try:
            arcpy.management.MakeFeatureLayer(target, sel)
            arcpy.management.SelectLayerByLocation(sel, "INTERSECT", join_layer)
            with arcpy.da.SearchCursor(sel, [id_field]) as cursor:
                for row in cursor:
                    for pid in _parse_ids(row[0], multi_id):
                        hits[pid] = {("Yes",)}
        finally:
            if arcpy.Exists(sel):
                arcpy.management.Delete(sel)

    return hits


def format_result(hits, pid, index=None, default="No"):
    """
    Render one pole's hits as a cell value.

    `index` picks a single field out of each tuple — that is how soil ends up as
    separate symbol and name columns. Without it, all fields are joined.
    `default` is what a structure with no hit prints, so a lookup that matched
    nothing can say so rather than printing a plausible-looking "No".
    """
    values = hits.get(pid)
    if not values:
        return default

    if index is None:
        rendered = {VALUE_SEP.join(v for v in tup if v) for tup in values}
    else:
        rendered = {tup[index] for tup in values if index < len(tup) and tup[index]}

    rendered = sorted(v for v in rendered if v)
    return ", ".join(rendered) if rendered else default


def check_layer(local, title):
    """Log what a layer brought to the project area, and whether it lines up with the poles."""
    count = feature_count(local)
    if count == 0:
        log("Warning", title,
            "0 features in the project area - every answer from this layer will be blank. "
            "Check that the service is reachable and covers the project.")
    else:
        log("Layer", title, f"{count} feature(s) in the project area")

    try:
        their_sr = arcpy.Describe(local).spatialReference
        if their_sr.factoryCode and pole_sr.factoryCode and their_sr.factoryCode != pole_sr.factoryCode:
            log("Layer", title,
                f"projected on the fly from {their_sr.name} to {pole_sr.name}")
    except Exception:
        pass

    return count


def analyze(layer, title, target, id_field, role=None, preferred=None, want=1, multi_id=False):
    """Resolve a layer's name field(s), then intersect it with `target`."""
    local = ensure_feature_layer(
        layer, gdb, aoi_features=aoi_features, aoi_extent_fc=aoi_extent_fc
    )
    count = check_layer(local, title)
    fields, provenance = resolve_fields(
        layer, title, local,
        role=role or infer_role(title),
        preferred=preferred, want=want, overrides=overrides,
    )
    log("Field", title, provenance)
    hits = get_intersect_hits(target, local, id_field, fields, gdb, multi_id=multi_id)
    return hits, local, count


def blank_answer(title, hits, count, expect_every_structure):
    """
    What to print when a structure has no hit, and a note when that looks wrong.

    "No" is a real answer for a constraint layer and a meaningless one for an
    administrative lookup — every structure sits in some county — so a lookup
    that matched nothing says so instead of quietly reading as a clean result.
    """
    matched = sum(1 for pid in pole_ids if hits.get(pid))
    log("Coverage", title, f"{matched} of {len(pole_ids)} structures matched")

    if count == 0:
        return "Layer empty"
    if expect_every_structure and matched == 0:
        log("Warning", title,
            "no structure matched this layer, which should not happen for a lookup that "
            "covers the whole state. Check the layer's extent, and the Run Log above for "
            "which field was read.")
        return "No match - see Run Log"
    return "No"


# ---------------------------------------------------------------------------
# Structure attributes
#
# The 2026 layer is a stream we cannot rename, and it has changed once already.
# Each attribute lists the 2026 name first, then older names, so a rename
# produces a warning and a blank column rather than a crash.
# ---------------------------------------------------------------------------

STRUCTURE_ATTRS = [
    ("Latitude", ["Latitude", "Structures_Latitude", "Lat"]),
    ("Longitude", ["Longitude", "Structures_Longitude", "Long", "Lon"]),
    ("Height", ["Height1", "Height", "Structures_Height1"]),
    ("Install_Date", ["Installati", "Installation Date", "Installation_Date", "InstallDate"]),
    ("Structure_Type", ["Structure T", "Structure_T", "Structure Type", "StructureType"]),
]

LEAD_COLUMNS = [
    "Pole_ID", "Latitude", "Longitude", "Height", "Install_Date", "Structure_Type",
    "County", "Township_City", "Sec_Twp_Rng", "Soil_Symbol", "Soil_Name", "Parcel",
    "MDOT_Jurisdiction", "AR_MDOT_Jurisdiction", "SESC_Trigger", "SESC_Features",
]

overrides = parse_overrides(field_override_param)
if overrides:
    log("Setup", "Field Overrides", f"{len(overrides)} layer(s) overridden")

# --- Buffer poles ---

pole_buffer_60 = os.path.join(gdb, "PoleBuffer_60ft")
arcpy.analysis.Buffer(poles, pole_buffer_60, "60 Feet", dissolve_option="NONE")

# The work zone: everything except SESC is a question about this buffer.
# The 500 ft buffer below is used only for the SESC check.
pole_buffer_500 = os.path.join(gdb, "PoleBuffer_500ft")
arcpy.analysis.Buffer(poles, pole_buffer_500, "500 Feet", dissolve_option="NONE")

# SESC is 500 ft from the structure by default. The work-zone option measures
# 500 ft from the edge of the 60 ft work zone instead, which is the stricter
# reading of "earth disturbance within 500 feet of a body of water".
if sesc_from_work_zone:
    sesc_target = os.path.join(gdb, "PoleBuffer_560ft")
    arcpy.analysis.Buffer(poles, sesc_target, "560 Feet", dissolve_option="NONE")
    log("Setup", "SESC", "measuring 560 ft (60 ft work zone + 500 ft)")
else:
    sesc_target = pole_buffer_500
    log("Setup", "SESC", "measuring 500 ft from the structure")

# --- Read structures ---

pole_fields = field_index(poles)
attr_fields = {}
for column, candidates in STRUCTURE_ATTRS:
    match = find_field(pole_fields, candidates)
    if match:
        attr_fields[column] = match
        log("Structure field", column, f"reading '{match}'")
    else:
        log("Structure field", column,
            f"not found (tried {', '.join(candidates)}); column left blank")

if not pole_id_field:
    arcpy.AddError("Pole ID Field is required.")
    raise arcpy.ExecuteError

pole_ids = []
pole_attrs = {}
duplicate_ids = set()

with arcpy.da.SearchCursor(poles, [pole_id_field] + list(attr_fields.values())) as cursor:
    for row in cursor:
        pid = "" if row[0] is None else str(row[0]).strip()
        if not pid:
            continue
        if pid in pole_attrs:
            duplicate_ids.add(pid)
            continue
        pole_ids.append(pid)
        pole_attrs[pid] = {col: row[i + 1] for i, col in enumerate(attr_fields)}

if duplicate_ids:
    shown = ", ".join(sorted(duplicate_ids)[:10])
    log("Warning", pole_id_field,
        f"{len(duplicate_ids)} duplicate ID(s) skipped after the first row: {shown}")

if not pole_ids:
    arcpy.AddError(
        f"No structures have a value in '{pole_id_field}'. Check the Pole ID Field parameter."
    )
    raise arcpy.ExecuteError

log("Setup", "Structures", f"{len(pole_ids)} structures to analyze")

# --- Build AOI for web layer downloads ---

aoi_buffer_dissolved = os.path.join(gdb, "aoi_buffer_dissolved")
arcpy.management.Dissolve(sesc_target, aoi_buffer_dissolved)

aoi_ar_dissolved = os.path.join(gdb, "aoi_ar_dissolved")
arcpy.management.Dissolve(access_routes, aoi_ar_dissolved)

aoi_features = [aoi_buffer_dissolved, aoi_ar_dissolved]

pole_ext = arcpy.Describe(sesc_target).extent
ar_ext = arcpy.Describe(access_routes).extent
sr = arcpy.Describe(sesc_target).spatialReference
pole_sr = sr

aoi_polygon = arcpy.Polygon(arcpy.Array([
    arcpy.Point(min(pole_ext.XMin, ar_ext.XMin), min(pole_ext.YMin, ar_ext.YMin)),
    arcpy.Point(min(pole_ext.XMin, ar_ext.XMin), max(pole_ext.YMax, ar_ext.YMax)),
    arcpy.Point(max(pole_ext.XMax, ar_ext.XMax), max(pole_ext.YMax, ar_ext.YMax)),
    arcpy.Point(max(pole_ext.XMax, ar_ext.XMax), min(pole_ext.YMin, ar_ext.YMin)),
]), sr)

aoi_extent_fc = os.path.join(gdb, "aoi_extent_poly")
arcpy.management.CopyFeatures([aoi_polygon], aoi_extent_fc)

# --- Seed results with the structure attributes ---

results = {}
for pid in pole_ids:
    row = {"Pole_ID": pid}
    row.update({col: pole_attrs[pid].get(col) for col in attr_fields})
    results[pid] = row

# --- Point-in-polygon lookups (the structure itself, not a buffer) ---

POINT_ANALYSES = [
    (county_layer, "County", ["County"], "county",
     [["LABEL", "NAME", "COUNTY"]], 1),
    (township_layer, "Township/City", ["Township_City"], "township",
     [["LABEL", "NAME"]], 1),
    (plss_layer, "Section-Township-Range", ["Sec_Twp_Rng"], "plss",
     [["STR_Label", "STR", "STRLABEL"]], 1),
    (soil_layer, "Soil", ["Soil_Symbol", "Soil_Name"], "soil",
     [["Mapunit Symbol", "MUSYM"], ["Mapunit Name", "MUNAME"]], 2),
    (parcel_layer, "Parcel", ["Parcel"], "parcel", None, 1),
]

for layer, title, columns, role, preferred, want in POINT_ANALYSES:
    if not layer:
        continue
    try:
        hits, _local, count = analyze(
            layer, title, poles, pole_id_field, role=role, preferred=preferred, want=want
        )
        # Parcels genuinely stop at the edge of a county's dataset; the rest are
        # statewide layers where a miss means something is wrong.
        default = blank_answer(title, hits, count, expect_every_structure=(role != "parcel"))
        for pid in pole_ids:
            for i, column in enumerate(columns):
                results[pid][column] = format_result(
                    hits, pid, index=i if len(columns) > 1 else None, default=default
                )
    except Exception as exc:
        log("Error", title, f"skipped - {exc}")
        for pid in pole_ids:
            for column in columns:
                results[pid][column] = "Layer failed - see Run Log"

# --- MDOT jurisdiction (work zone and access routes, since both enter the ROW) ---

if mdot_layer:
    try:
        local = ensure_feature_layer(
            mdot_layer, gdb, aoi_features=aoi_features, aoi_extent_fc=aoi_extent_fc
        )
        fields, provenance = resolve_fields(
            mdot_layer, "MDOT Jurisdiction", local, role="road",
            preferred=[["Legal System", "LegalSystem", "LEGALSYST"]], overrides=overrides,
        )
        log("Field", "MDOT Jurisdiction", provenance)

        count = check_layer(local, "MDOT Jurisdiction")
        hits_wz = get_intersect_hits(pole_buffer_60, local, pole_id_field, fields, gdb)
        hits_ar = get_intersect_hits(
            access_routes, local, ar_pole_id_field, fields, gdb, multi_id=True
        )
        default = blank_answer(
            "MDOT Jurisdiction", hits_wz, count, expect_every_structure=False
        )
        for pid in pole_ids:
            results[pid]["MDOT_Jurisdiction"] = format_result(hits_wz, pid, default=default)
            results[pid]["AR_MDOT_Jurisdiction"] = format_result(hits_ar, pid, default=default)
    except Exception as exc:
        log("Error", "MDOT Jurisdiction", f"skipped - {exc}")
        for pid in pole_ids:
            results[pid]["MDOT_Jurisdiction"] = "Layer failed - see Run Log"
            results[pid]["AR_MDOT_Jurisdiction"] = "Layer failed - see Run Log"

# --- SESC ---
#
# Which layers count as a "body of water" is a county decision, so the layers
# are chosen at run time rather than hardcoded. Wetland layers also contribute
# their type field, so PEM and PFO are visible in the answer.

if sesc_layers:
    sesc_hits = {pid: set() for pid in pole_ids}
    sesc_empty = []
    sesc_failed = []
    for layer in sesc_layers:
        title = layer_title(layer)
        try:
            local = ensure_feature_layer(
                layer, gdb, aoi_features=aoi_features, aoi_extent_fc=aoi_extent_fc
            )
            fields, provenance = resolve_fields(
                layer, title, local, role=infer_role(title), overrides=overrides
            )
            log("Field", f"SESC / {title}", provenance)

            if check_layer(local, f"SESC / {title}") == 0:
                sesc_empty.append(title)

            type_field = find_field(
                field_index(local),
                ["ATTRIBUTE", "WETLAND_TYPE", "WETLAND_TY", "WETLAND_CLASS", "CLASS"],
            )
            if type_field and type_field not in fields:
                fields = fields + [type_field]

            hits = get_intersect_hits(sesc_target, local, pole_id_field, fields, gdb)
            for pid, values in hits.items():
                if pid in sesc_hits:
                    sesc_hits[pid] |= values
        except Exception as exc:
            log("Error", f"SESC / {title}", f"skipped - {exc}")
            sesc_failed.append(title)

    # A clean SESC result is only trustworthy if every chosen layer was readable.
    unusable = sesc_empty + sesc_failed
    if unusable:
        log("Warning", "SESC",
            f"{len(unusable)} of {len(sesc_layers)} layer(s) returned nothing "
            f"({', '.join(unusable)}); a 'No' below is not a clean result")
        clean = f"Incomplete - {', '.join(unusable)} unreadable"
    else:
        clean = "No"

    for pid in pole_ids:
        results[pid]["SESC_Trigger"] = "Yes" if sesc_hits[pid] else clean
        results[pid]["SESC_Features"] = format_result(sesc_hits, pid, default=clean)

# --- Constraint layers ---

for constraint in constraint_layers:
    constraint_name = layer_title(constraint)
    try:
        local = ensure_feature_layer(
            constraint, gdb, aoi_features=aoi_features, aoi_extent_fc=aoi_extent_fc
        )
        fields, provenance = resolve_fields(
            constraint, constraint_name, local,
            role=infer_role(constraint_name), overrides=overrides,
        )
        log("Field", constraint_name, provenance)

        count = check_layer(local, constraint_name)
        hits_60 = get_intersect_hits(pole_buffer_60, local, pole_id_field, fields, gdb)
        hits_ar = get_intersect_hits(
            access_routes, local, ar_pole_id_field, fields, gdb, multi_id=True
        )

        default = blank_answer(
            constraint_name, hits_60, count, expect_every_structure=False
        )
        for pid in pole_ids:
            results[pid][f"60ft_Pole_{constraint_name}"] = format_result(
                hits_60, pid, default=default
            )
            results[pid][f"AR_{constraint_name}"] = format_result(
                hits_ar, pid, default=default
            )
    except Exception as exc:
        log("Error", constraint_name, f"skipped - {exc}")
        for pid in pole_ids:
            results[pid][f"60ft_Pole_{constraint_name}"] = "Layer failed - see Run Log"
            results[pid][f"AR_{constraint_name}"] = "Layer failed - see Run Log"

# --- Two mile T&E review ---
#
# Project-wide rather than per-pole: one buffer around the whole footprint
# covers a linear project that a single circle drawn from a central structure
# would miss, and costs one join instead of one per structure.

te_df = None
if run_te_review and mnfi_layer:
    try:
        te_poles = os.path.join(gdb, "te_buffer_poles")
        arcpy.analysis.Buffer(poles, te_poles, "2 Miles", dissolve_option="ALL")

        te_routes = os.path.join(gdb, "te_buffer_routes")
        arcpy.analysis.Buffer(access_routes, te_routes, "2 Miles", dissolve_option="ALL")

        te_merged = os.path.join(gdb, "te_buffer_merged")
        arcpy.management.Merge([te_poles, te_routes], te_merged)

        te_buffer = os.path.join(gdb, "te_buffer_2mi")
        arcpy.management.Dissolve(te_merged, te_buffer)

        mnfi_local = ensure_feature_layer(
            mnfi_layer, gdb, aoi_features=[te_buffer], aoi_extent_fc=te_buffer
        )

        sel = temp_name("mnfi_sel")
        try:
            arcpy.management.MakeFeatureLayer(mnfi_local, sel)
            arcpy.management.SelectLayerByLocation(sel, "INTERSECT", te_buffer)

            mnfi_fields = field_index(mnfi_local)
            names = [name for name, _a, _t in mnfi_fields]
            headers = [alias for _n, alias, _t in mnfi_fields]

            rows = []
            with arcpy.da.SearchCursor(sel, names) as cursor:
                for row in cursor:
                    rows.append(tuple("" if v is None else v for v in row))
        finally:
            if arcpy.Exists(sel):
                arcpy.management.Delete(sel)

        te_df = pd.DataFrame(rows, columns=headers).drop_duplicates()
        log("T&E", "MNFI", f"{len(te_df)} occurrence(s) within 2 miles of the project")
    except Exception as exc:
        log("Error", "T&E 2 mile review", f"skipped - {exc}")
elif run_te_review:
    log("Warning", "T&E 2 mile review", "requested but no MNFI layer supplied; skipped")

# --- Assemble the matrix ---

df = pd.DataFrame.from_dict(results, orient="index")

if "Install_Date" in df.columns:
    df["Install_Date"] = (
        pd.to_datetime(df["Install_Date"], errors="coerce").dt.strftime("%Y-%m-%d").fillna("")
    )

lead = [c for c in LEAD_COLUMNS if c in df.columns]
df = df[lead + [c for c in df.columns if c not in lead]].reset_index(drop=True)

# --- Export ---

excel_path = os.path.join(out_folder, f"{project_name}_Pole_Constraint_Matrix.xlsx")

with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name="Pole Matrix", index=False)
    if te_df is not None:
        te_df.to_excel(writer, sheet_name="T&E 2mi", index=False)
    pd.DataFrame(RUN_LOG).to_excel(writer, sheet_name="Run Log", index=False)


def is_hit_column(header):
    """True for columns that report an overlap, as opposed to plain information."""
    if not header:
        return False
    if header.startswith("60ft_Pole_") or header.startswith("AR_"):
        return True
    return header in ("MDOT_Jurisdiction", "SESC_Trigger", "SESC_Features")


# Answers that mean "the tool could not tell you", which must not be mistaken
# for a clean result at a glance.
PROBLEM_ANSWERS = ("Layer empty", "Layer failed", "No match", "Incomplete")


orange_fill = PatternFill(start_color="FFCC99", end_color="FFCC99", fill_type="solid")
grey_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
black_font = Font(color="000000", bold=True)
problem_font = Font(color="9C0006", italic=True)

wb = load_workbook(excel_path)
ws = wb["Pole Matrix"]

headers = {cell.column: cell.value for cell in ws[1]}

for row in ws.iter_rows(min_row=2):
    for cell in row:
        val = str(cell.value) if cell.value else ""
        if not val:
            continue

        # Checked on every column, because a county or soil lookup that could
        # not be answered lands in an otherwise informational column.
        if val.startswith(PROBLEM_ANSWERS):
            cell.fill = grey_fill
            cell.font = problem_font
            continue

        if val != "No" and is_hit_column(headers.get(cell.column, "")):
            cell.fill = orange_fill
            cell.font = black_font

# Readable widths — the name columns are much wider than the Yes/No they replaced.
for sheet in wb.worksheets:
    sheet.freeze_panes = "A2"
    for column_cells in sheet.columns:
        longest = max(
            (len(str(c.value)) for c in column_cells[:200] if c.value is not None),
            default=10,
        )
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(longest + 2, 10), 45)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

wb.save(excel_path)
arcpy.AddMessage(f"Matrix created: {excel_path}")

# --- Add buffer layers to Contents pane ---

arcpy.SetParameter(18, pole_buffer_60)
arcpy.SetParameter(19, pole_buffer_500)
